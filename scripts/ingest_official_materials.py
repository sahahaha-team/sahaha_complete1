"""Index the supplied reports and the *current pages* listed in the Q&A workbook.

The workbook answers are deliberately never indexed: they are proposed answers,
not proof that a fee, benefit, or office hours are still current.
"""

from __future__ import annotations

import argparse
import hashlib
import logging
import re
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from crawler.saha_crawler import SahaCrawler
from database_db.database import Database
from database_db.vector_store import VectorStore
from processor.data_cleaner import CleanedChunk, DataCleaner
from processor.metadata_defaults import fallback_keywords

logger = logging.getLogger(__name__)
DEFAULT_DIR = Path("data/official_sources")
REPORTS = {
    "2026년 정보공개청구 민원 데이터 분석 결과보고.pdf": {
        "issued": "2026-08-04", "period": "2023-01-01~2025-12-31", "category": "정보공개",
    },
    "사하구 홈페이지 방문자 및 검색 통계 분석 결과 보고.pdf": {
        "issued": "2025-08-19", "period": "2023-01-01~2024-12-31 (메뉴별 방문: 2024-11-01~2025-04-30)",
        "category": "사하소개",
    },
}


def _official_url(url: str) -> bool:
    parsed = urlparse(url)
    return parsed.scheme == "https" and (parsed.hostname == "saha.go.kr" or
        (parsed.hostname or "").endswith(".saha.go.kr"))


def workbook_targets(path: Path) -> list[tuple[str, str]]:
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book["예상질문답변"]
    targets = []
    seen = set()
    for _number, question, _proposed_answer, url in sheet.iter_rows(min_row=2, values_only=True):
        url = str(url or "").strip()
        if not question or not _official_url(url):
            logger.warning("공식 URL이 없는 질문 제외: %s", question)
            continue
        if url not in seen:
            targets.append((str(question).strip(), url))
            seen.add(url)
    book.close()
    return targets


def _fetch_page(target: tuple[str, str]):
    question, url = target
    try:
        response = requests.get(url, timeout=20, headers={"User-Agent": "SahaOfficialKnowledgeRefresh/1.0"})
        response.raise_for_status()
        if not _official_url(response.url):
            raise ValueError("공식 도메인 밖으로 이동")
        response.encoding = response.apparent_encoding or "utf-8"
        soup = BeautifulSoup(response.text, "lxml")
        crawler = object.__new__(SahaCrawler)
        content = crawler._extract_content(soup)
        title = crawler._extract_title(soup)
        if len(content.strip()) < 80:
            raise ValueError("본문이 너무 짧음")
        return {"question": question, "url": url, "title": title,
                "content": content, "checked_at": datetime.now(timezone.utc).isoformat()}
    except Exception as exc:
        logger.warning("공식 페이지 갱신 실패 (%s): %s", url, exc)
        return None


def fetch_pages(targets: list[tuple[str, str]]) -> list[dict]:
    pages = []
    with ThreadPoolExecutor(max_workers=3) as pool:
        futures = [pool.submit(_fetch_page, target) for target in targets]
        for future in as_completed(futures):
            row = future.result()
            if row:
                pages.append(row)
    return sorted(pages, key=lambda row: row["url"])


def report_pages(path: Path) -> list[dict]:
    import pdfplumber

    config = REPORTS[path.name]
    output = []
    with pdfplumber.open(path) as pdf:
        # Page 1 is a cover summary. The 2026 cover's years conflict with the
        # methodology on page 2, so use the detailed report only.
        for page_number, pdf_page in enumerate(pdf.pages[1:], start=2):
            content = pdf_page.extract_text() or ""
            content = re.sub(r"^저장 :[^\n]*\n", "", content)
            content = re.sub(r"^문서관리카드[^\n]*\n", "", content)
            if len(content.strip()) < 80:
                logger.warning("보고서 %s %s쪽 텍스트 없음", path.name, page_number)
                continue
            digest = hashlib.sha256(path.name.encode()).hexdigest()[:16]
            output.append({
                "url": f"file://official_reports/{digest}/page-{page_number}",
                "title": f"{path.stem} ({page_number}쪽)",
                "content": f"자료 기준 기간: {config['period']}\n보고서 작성일: {config['issued']}\n{content}",
                "category": config["category"], "source_type": "official_report",
                "issued_at": config["issued"], "data_period": config["period"],
                "page_number": page_number, "source_title": path.stem,
            })
    return output


def _category(url: str) -> str:
    parsed = urlparse(url)
    if parsed.path.startswith("/health/"):
        return "보건소"
    prefix = parse_qs(parsed.query).get("mId", [""])[0][:2]
    return {"01": "전자민원", "03": "정보공개", "04": "분야별정보",
            "05": "사하복지", "06": "사하소개"}.get(prefix, "공식홈페이지")


def _service_type(category: str) -> str:
    return {"전자민원": "민원", "사하복지": "복지", "보건소": "보건",
            "정보공개": "기타"}.get(category, "기타")


def _make_chunks(cleaner: DataCleaner, page) -> list[CleanedChunk]:
    """Keep distinct benefit programs apart when one web page lists many amounts."""
    if "mId=0510070100" not in page.url:
        return cleaner.process(page)
    text = page.content
    markers = ("첫만남이용권", "출산지원금(구비, 시비)", "출산지원금(장애인)")
    positions = [text.find(marker) for marker in markers]
    if any(position < 0 for position in positions) or positions != sorted(positions):
        raise ValueError("출산 지원 페이지의 항목 구조가 바뀌어 자동 분리할 수 없음")
    parts = [text[positions[0]:positions[1]], text[positions[1]:positions[2]], text[positions[2]:]]
    text_chunks = []
    for heading, part in zip(markers, parts):
        cleaned = cleaner.clean_text(part)
        for piece in cleaner.splitter.split_text(cleaned):
            text_chunks.append(f"{heading}\n{piece}")
    return [CleanedChunk(
        chunk_id=hashlib.md5(f"{page.url}_{index}".encode()).hexdigest(),
        url=page.url, title=page.title, content=content,
        category=page.category, sub_category=page.sub_category,
        chunk_index=index, total_chunks=len(text_chunks),
    ) for index, content in enumerate(text_chunks)]


def _index_page(db: Database, vs: VectorStore, row: dict, *, force: bool = False) -> int:
    is_report = row["source_type"] == "official_report"
    page = type("Page", (), {})()
    page.url = row["url"]
    page.title = row["title"]
    page.content = row["content"]
    page.category = row["category"]
    page.sub_category = "공식 분석 보고서" if is_report else "100개 질문 관련 공식 페이지"
    page.attachments = []

    status = db.upsert_raw_page(page)
    if not is_report:
        db.mark_page_checked(page.url)
    existing = db.client.table("documents").select("id,metadata").eq("metadata->>url", page.url).limit(1).execute().data or []
    content_hash = hashlib.sha256(page.content.encode()).hexdigest()
    already_indexed = bool(existing and
        (existing[0].get("metadata") or {}).get("source_type") == row["source_type"] and
        (existing[0].get("metadata") or {}).get("service_type") ==
            ("통계" if is_report else _service_type(page.category)) and
        (existing[0].get("metadata") or {}).get("content_hash") == content_hash)
    if status == "unchanged" and already_indexed and not force:
        return 0

    cleaner = DataCleaner()
    chunks = _make_chunks(cleaner, page)
    if not chunks:
        raise ValueError(f"청크 생성 실패: {page.url}")
    tagged = []
    pairs = []
    for chunk in chunks:
        metadata = {"service_type": "통계" if is_report else _service_type(page.category), "department": "",
                    "keywords": fallback_keywords(chunk.content, page.title), "summary": ""}
        tagged.append((chunk, metadata))
        vector_meta = {
            "url": page.url, "title": page.title, "category": page.category,
            "sub_category": page.sub_category, "service_type": metadata["service_type"],
            "department": "", "keywords": metadata["keywords"],
            "source_type": row["source_type"], "content_hash": content_hash,
        }
        if is_report:
            vector_meta.update({key: row[key] for key in
                ("issued_at", "data_period", "page_number", "source_title")})
        else:
            vector_meta["verified_at"] = row["checked_at"]
        pairs.append((chunk, vector_meta))

    db.save_chunks_bulk(tagged)
    vs.add_chunks_batch(pairs, batch_size=30, db=db)
    valid_ids = {chunk.chunk_id for chunk in chunks}
    db.delete_stale_derived_for_urls({page.url: valid_ids})
    # Remove old vectors for this URL even if they were made by another import.
    all_docs = db.client.table("documents").select("id").eq("metadata->>url", page.url).execute().data or []
    old_ids = [item["id"] for item in all_docs if item["id"] not in valid_ids]
    for start in range(0, len(old_ids), 100):
        db.client.table("documents").delete().in_("id", old_ids[start:start + 100]).execute()
    return len(chunks)


def ingest(source_dir: Path, *, reports: bool = True, pages: bool = True,
           force: bool = False, only_url: str | None = None) -> dict:
    rows = []
    if reports:
        for filename in REPORTS:
            path = source_dir / filename
            if not path.is_file():
                raise FileNotFoundError(path)
            rows.extend(report_pages(path))
    fetched = []
    failed_urls = []
    if pages:
        workbook = source_dir / "사하구 홈페이지_100개 질문.xlsx"
        if not workbook.is_file():
            raise FileNotFoundError(workbook)
        targets = workbook_targets(workbook)
        if only_url:
            targets = [target for target in targets if target[1] == only_url]
            if not targets:
                raise ValueError("엑셀에 해당 공식 URL이 없습니다")
        fetched = fetch_pages(targets)
        failed_urls = sorted(set(url for _question, url in targets) - {row["url"] for row in fetched})
        if not fetched:
            raise RuntimeError("공식 페이지를 한 건도 확인하지 못했습니다")
        for row in fetched:
            row.update({"category": _category(row["url"]), "source_type": "verified_page"})
        rows.extend(fetched)
        logger.info("공식 페이지 확인: %d/%d", len(fetched), len(targets))

    db = Database(admin=True)
    if failed_urls:
        # A failed verification must not leave an old page eligible as "current".
        db.client.table("raw_pages").update({
            "last_checked_at": "1970-01-01T00:00:00+00:00"
        }).in_("url", failed_urls).execute()
    vs = VectorStore(admin=True)
    chunks = 0
    errors = 0
    for i, row in enumerate(rows, 1):
        try:
            chunks += _index_page(db, vs, row, force=force)
            logger.info("적재 %d/%d: %s", i, len(rows), row["title"])
        except Exception:
            errors += 1
            logger.exception("적재 실패: %s", row["url"])
    if rows:
        from chatbot.bm25_index import BM25Index
        BM25Index().rebuild()
    return {"reports": len(rows) - len(fetched), "verified_pages": len(fetched),
            "failed_pages": len(failed_urls), "new_chunks": chunks, "errors": errors}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_DIR)
    parser.add_argument("--reports-only", action="store_true")
    parser.add_argument("--pages-only", action="store_true")
    parser.add_argument("--force", action="store_true", help="Rebuild vectors even when text is unchanged")
    parser.add_argument("--url", help="Refresh one official URL from the workbook")
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    logging.getLogger("httpx").setLevel(logging.WARNING)
    print(ingest(args.source_dir, reports=not args.pages_only, pages=not args.reports_only,
                 force=args.force, only_url=args.url))
