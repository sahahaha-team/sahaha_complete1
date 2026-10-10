"""구청 제공 예상질문 엑셀을 버전 관리 가능한 공식 FAQ JSON으로 변환한다."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


TIME_SENSITIVE_PATTERN = re.compile(
    r"(?:20\d{2}년|\d{1,2}:\d{2}|월\s*\d+만원|최대\s*\d|\d+일\s*이내|"
    r"운영시간|진료시간|접수시간|수수료|비용|지원금|금액|공시)"
)

CATEGORY_RULES = [
    ("보건·의료", ("보건", "건강", "접종", "검진", "의료비", "치매", "결핵", "임산부", "암")),
    ("환경·폐기물", ("쓰레기", "폐기물", "정화조", "환경", "오존", "석면", "탄소", "멧돼지")),
    ("복지·돌봄", ("복지", "수급", "연금", "돌봄", "출산", "1인가구", "장애인", "청소년부모")),
    ("세금·부동산", ("지방세", "재산세", "자동차세", "취득세", "공시지가", "주택가격", "부동산", "임대주택", "공동주택", "조상 땅")),
    ("민원·여권", ("민원", "여권", "정부24", "발급", "수수료")),
    ("교통·안전", ("주차", "버스", "도시철도", "자전거", "대피", "지진", "폭염", "보험")),
    ("교육·일자리", ("교육", "학습", "구직", "근로", "일자리", "토익")),
]


def category_for(question: str, answer: str) -> str:
    text = f"{question} {answer}"
    for category, keywords in CATEGORY_RULES:
        if any(keyword in text for keyword in keywords):
            return category
    return "기타"


def keywords_for(question: str, limit: int = 8) -> list[str]:
    stopwords = {"사하구", "사하구청", "어떻게", "어디", "무엇", "있나요", "인가요", "받을", "대한"}
    unique = []
    for token in re.findall(r"[가-힣A-Za-z0-9]{2,}", question or ""):
        if token not in stopwords and token not in unique:
            unique.append(token)
    return unique[:limit]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, default=Path("resources/official_faq.json"))
    parser.add_argument("--verified-at", default="2026-10-06")
    args = parser.parse_args()

    workbook = load_workbook(args.source, read_only=True, data_only=True)
    sheet = workbook["예상질문답변"]
    items = []
    for number, question, answer, url in sheet.iter_rows(min_row=2, values_only=True):
        if not question:
            continue
        if not answer or not url:
            raise ValueError(f"행 {number}: 예상답변 또는 위치가 없습니다.")
        items.append({
            "id": int(number),
            "question": str(question).strip(),
            "answer": str(answer).strip(),
            "url": str(url).strip(),
            "category": category_for(str(question), str(answer)),
            "keywords": keywords_for(str(question)),
            "is_time_sensitive": bool(TIME_SENSITIVE_PATTERN.search(str(answer))),
            "verified_at": args.verified_at,
        })
    workbook.close()

    if len(items) != 100:
        raise ValueError(f"공식 예상질문은 100개여야 합니다. 현재 {len(items)}개")
    if len({item["url"] for item in items}) != 100:
        raise ValueError("공식 예상질문의 URL은 100개가 모두 고유해야 합니다.")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(items, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"saved={len(items)} output={args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
