"""Use the supplied question list to select an official page, never its draft answer."""

from __future__ import annotations

import re
import json
from difflib import SequenceMatcher
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook
from chatbot.question_intent import asks_opening_hours, asks_location
from chatbot.query_subject import fallback_keywords
from chatbot.vaccination import vaccine_place_page
from chatbot.official_faq import OfficialFAQIndex, PRIORITY_SERVICE_PATH

WORKBOOK = Path(__file__).resolve().parents[1] / "data" / "official_sources" / "사하구 홈페이지_100개 질문.xlsx"


def current_official_url(url: str) -> str:
    path = Path(__file__).resolve().parents[1] / 'resources' / 'official_source_redirects.json'
    replacements = json.loads(path.read_text(encoding='utf-8')) if path.is_file() else {}
    replacement = replacements.get(url, url)
    from chatbot.reference_answers import safe_official_link
    return replacement if safe_official_link(replacement) else url


@lru_cache(maxsize=1)
def _question_indexes() -> tuple[OfficialFAQIndex, OfficialFAQIndex]:
    return OfficialFAQIndex(), OfficialFAQIndex(PRIORITY_SERVICE_PATH)


def priority_entities(question: str) -> list[str]:
    """Distinctive registered names are URL hints, never answer facts."""
    items = _question_indexes()[1].items
    entities = {str(term) for item in items for term in item.get('navigation_terms', [])
                if len(_normalize(str(term))) >= 3}
    normalized = _normalize(question)
    return sorted((name for name in entities if _normalize(name) in normalized), key=lambda name: (-len(name), name))


def _priority_entity_url(question: str) -> str | None:
    names = priority_entities(question)
    from chatbot.query_subject import service_names
    urls = {current_official_url(item['url']) for item in _question_indexes()[1].items
            if any(name in item.get('navigation_terms', []) for name in names)
            and service_names(question) <= service_names(item['question'])}
    return next(iter(urls)) if len(urls) == 1 else None


def priority_source_name(url: str) -> str:
    return next((str(item.get('source_name') or '') for item in _question_indexes()[1].items
                 if current_official_url(item['url']) == url), '')


def find_official_question(question: str) -> dict | None:
    """FAQ similarity is a navigation hint, never proof of a draft answer."""
    matches = [index.find_best(question, min_score=0.88) for index in _question_indexes()]
    match = max((match for match in matches if match), key=lambda match: match['score'], default=None)
    return {**match, 'url': current_official_url(match['url'])} if match else None


def _normalize(question: str) -> str:
    return re.sub(r"[^가-힣a-z0-9]", "", (question or "").lower())


@lru_cache(maxsize=1)
def _targets() -> list[tuple[str, str]]:
    # Committed JSON keeps the official question routes usable on a new PC.
    output = [(_normalize(str(item['question'])), current_official_url(str(item['url'])))
              for index in _question_indexes() for item in index.items
              if str(item.get('url', '')).startswith('https://www.saha.go.kr/')]
    if not WORKBOOK.is_file():
        return output
    book = load_workbook(WORKBOOK, read_only=True, data_only=True)
    try:
        for _number, question, _draft_answer, url in book["예상질문답변"].iter_rows(min_row=2, values_only=True):
            if question and url and str(url).startswith("https://www.saha.go.kr/"):
                output.append((_normalize(str(question)), current_official_url(str(url))))
    finally:
        book.close()
    return output


def confident_official_page(question: str) -> str | None:
    """A registered question identifies a business, never its factual answer."""
    normalized = _normalize(question)
    exact = {url for candidate, url in _targets() if candidate == normalized}
    if exact:
        return next(iter(exact)) if len(exact) == 1 else None
    if len(normalized) < 12:
        return None
    scores = [(SequenceMatcher(None, normalized, candidate).ratio(), url) for candidate, url in _targets()]
    best = max((score for score, _ in scores), default=0)
    urls = {url for score, url in scores if score >= max(.92, best - .01)}
    return next(iter(urls)) if len(urls) == 1 else None


def match_official_page(question: str) -> str | None:
    normalized = _normalize(question)
    vaccine_url = vaccine_place_page(question)
    if vaccine_url:
        return vaccine_url
    named_url = _priority_entity_url(question)
    if named_url:
        return named_url
    from chatbot.official_faq import required_faq_ids
    if required_faq_ids(question) == {1, 2}:
        return 'https://www.saha.go.kr/portal/contents.do?mId=0604050000'
    if asks_location(question) and fallback_keywords(question) == {'보건소'}:
        return 'https://www.saha.go.kr/health/contents.do?mId=0103000000'
    if asks_location(question) and fallback_keywords(question) == {'구청'}:
        return 'https://www.saha.go.kr/portal/contents.do?mId=0604050000'
    # The health-center directions page also contains its general clinical
    # hours. Do not substitute a village health center or a service's hours.
    if ('보건소' in normalized and asks_opening_hours(question)
            and not any(word in normalized for word in (
                '마을건강', '건강생활지원', '보건지소', '치매', '금연', '접종',
                '검사', '보건증', '건강진단', '검진', '모자보건', '구강', '물리치료',
            ))):
        return 'https://www.saha.go.kr/health/contents.do?mId=0103000000'
    # Confirmed service + question goal selects its verified guide directly.
    # Conversational qualifiers like '성인이야/처음이야' must not let an
    # unrelated adult vaccination page crowd the passport guide out of top 5.
    if '여권' in normalized and any(word in normalized for word in ('준비', '서류', '구비', '사진', '발급기간', '기간은', '며칠', '소요')):
        return 'https://www.saha.go.kr/portal/contents.do?mId=0104020000'
    if any(word in normalized for word in ('불법주차', '불법주정차')) and '신고' in normalized:
        return 'https://www.saha.go.kr/portal/contents.do?mId=0403080000'
    if '대형폐기물' in normalized and any(word in normalized for word in ('배출', '방법', '접수', '수수료')):
        return 'https://www.saha.go.kr/portal/contents.do?mId=0405050103'
    if len(normalized) < 7:
        return None
    targets = _targets()
    for candidate, url in targets:
        if normalized == candidate:
            return url
    match = find_official_question(question)
    if match and str(match.get('url', '')).startswith('https://www.saha.go.kr/'):
        return match['url']
    # Minor spelling/wording changes are allowed; unrelated topics stay in RAG.
    close = [(SequenceMatcher(None, normalized, candidate).ratio(), url) for candidate, url in targets]
    if not close:
        return None
    score, url = max(close)
    return url if score >= 0.92 else None
